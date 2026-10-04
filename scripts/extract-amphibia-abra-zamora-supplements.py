from pathlib import Path
import hashlib,json,zipfile,xml.etree.ElementTree as ET
import openpyxl
root=Path('data/sources/snapshots/amphibia-primary-corpus-2026-10-02/PMC7482940')
result={'pmc':'PMC7482940','rawFilesUnmodified':True,'biologicalAssignmentsCertified':False,'documents':[],'workbooks':[]}
for path in sorted(root.glob('pone.0238306.s00*')):
    before=hashlib.sha256(path.read_bytes()).hexdigest()
    if path.suffix=='.docx':
        with zipfile.ZipFile(path) as archive:
            doc=ET.fromstring(archive.read('word/document.xml'))
        ns={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
        paragraphs=[''.join(p.itertext()) for p in doc.findall('.//w:body/w:p',ns)]
        tables=[]
        for ti,t in enumerate(doc.findall('.//w:tbl',ns),1):
            rows=[[''.join(x.itertext()) for x in row.findall('w:tc',ns)] for row in t.findall('w:tr',ns)]
            tables.append({'table':ti,'rows':[{'row':i,'cells':r} for i,r in enumerate(rows,1)]})
        result['documents'].append({'file':path.name,'sha256':before,'paragraphs':paragraphs,'tables':tables})
    elif path.suffix=='.xlsx':
        wb=openpyxl.load_workbook(path,read_only=True,data_only=False)
        sheets=[]
        for ws in wb.worksheets:
            rows=[]
            for row_number,row in enumerate(ws.iter_rows(),1):
                cells=[{'cell':c.coordinate,'value':c.value,'dataType':c.data_type} for c in row if c.value is not None]
                if cells:rows.append({'row':row_number,'cells':cells})
            sheets.append({'sheet':ws.title,'maxRow':ws.max_row,'maxColumn':ws.max_column,'rows':rows})
        wb.close();result['workbooks'].append({'file':path.name,'sha256':before,'sheets':sheets})
    if before!=hashlib.sha256(path.read_bytes()).hexdigest():raise RuntimeError('Raw source changed')
output=root/'abra-zamora-supplement-extraction.json'
output.write_text(json.dumps(result,ensure_ascii=False,indent=2,default=str)+'\n',encoding='utf8')
print(json.dumps({'documents':[{'file':x['file'],'paragraphs':len(x['paragraphs']),'tables':[len(t['rows']) for t in x['tables']]} for x in result['documents']],'workbooks':[{'file':x['file'],'sheets':[{'name':s['sheet'],'rows':s['maxRow'],'cols':s['maxColumn']} for s in x['sheets']]} for x in result['workbooks']],'biologicalAssignmentsCertified':False}))
